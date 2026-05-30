"""Best-effort text extraction for the engine's content stage.

Each backend is optional and degrades gracefully: when neither a system tool
nor a Python library can read a format, extraction returns ``""`` and the engine
simply falls back to filename / type rules. Nothing here raises to the caller,
so a corrupt or password-protected file never breaks a sort run.
"""

from __future__ import annotations

import html
import re
import shutil
import subprocess
from pathlib import Path

# Formats we know how to read as text. Anything else skips the content stage.
PARSABLE: frozenset[str] = frozenset(
    {
        "pdf",
        "docx",
        "doc",
        "rtf",
        "html",
        "htm",
        "txt",
        "md",
        "csv",
        "tsv",
        "log",
        "xlsx",
    }
)

_TAG_RE = re.compile(r"<[^>]+>")
_TIMEOUT = 15


def _run(cmd: list[str]) -> str:
    try:
        result = subprocess.run(cmd, capture_output=True, timeout=_TIMEOUT)
        return result.stdout.decode("utf-8", "ignore")
    except Exception:
        return ""


def _from_pdf(path: Path) -> str:
    if shutil.which("pdftotext"):  # poppler, common on macOS and Linux
        text = _run(["pdftotext", "-l", "2", "-q", str(path), "-"])
        if text.strip():
            return text
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = reader.pages[:2]
        return "\n".join((page.extract_text() or "") for page in pages)
    except Exception:
        return ""


def _from_docx(path: Path) -> str:
    try:
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    except Exception:
        pass
    if shutil.which("textutil"):  # macOS native
        return _run(["textutil", "-convert", "txt", "-stdout", str(path)])
    return ""


def _from_legacy_office(path: Path) -> str:
    if shutil.which("textutil"):  # macOS reads .doc/.rtf natively
        return _run(["textutil", "-convert", "txt", "-stdout", str(path)])
    for tool in ("antiword", "catdoc"):  # common on Linux
        if shutil.which(tool):
            return _run([tool, str(path)])
    return ""


def _from_html(path: Path) -> str:
    if shutil.which("textutil"):
        text = _run(["textutil", "-convert", "txt", "-stdout", str(path)])
        if text.strip():
            return text
    try:
        raw = path.read_text("utf-8", "ignore")
        return html.unescape(_TAG_RE.sub(" ", raw))
    except Exception:
        return ""


def _from_text(path: Path, max_bytes: int) -> str:
    try:
        with path.open("r", encoding="utf-8", errors="ignore") as handle:
            return handle.read(max_bytes)
    except Exception:
        return ""


def _from_xlsx(path: Path) -> str:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        cells: list[str] = []
        for row in sheet.iter_rows(max_row=20, values_only=True):
            cells += [str(c) for c in row if c is not None]
        return " ".join(cells)
    except Exception:
        return ""


def extract_text(path: Path, ext: str, max_bytes: int = 4000) -> str:
    """Return up to ``max_bytes`` of extracted text, or ``""`` if unsupported."""
    if not path.is_file():
        return ""
    ext = ext.lower()
    if ext == "pdf":
        text = _from_pdf(path)
    elif ext == "docx":
        text = _from_docx(path)
    elif ext in {"doc", "rtf"}:
        text = _from_legacy_office(path)
    elif ext in {"html", "htm"}:
        text = _from_html(path)
    elif ext in {"txt", "md", "csv", "tsv", "log"}:
        text = _from_text(path, max_bytes)
    elif ext == "xlsx":
        text = _from_xlsx(path)
    else:
        text = ""
    return text[:max_bytes]

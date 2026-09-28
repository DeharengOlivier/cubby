"""The Python document parsers, run in a child process by the extraction adapter.

``pypdf``, ``python-docx`` and ``openpyxl`` parse whatever lands in the
Downloads folder. A crafted file can make a parser loop, or inflate a zip into
gigabytes. Inside the agent's own process that would stall or kill the agent;
in a child it costs one timeout. So these functions run in the parent only in
tests, and in production through::

    python -m cubby.adapters.parsers <pdf|docx|xlsx> <path> <max_chars>

which prints the text on stdout. The parent sets the timeout, the memory
ceiling and a neutral working directory (see ``extraction._run``). A parser
that raises (a corrupt file, a ``MemoryError`` under the ceiling) makes the
child exit with :data:`EXIT_PARSE_FAILED` and name the exception on stderr,
so the parent can report it; the engine still falls back to the filename and
type stages, as designed.
"""

from __future__ import annotations

import sys


def pdf_text(path: str, max_chars: int) -> str:
    from pypdf import PdfReader  # noqa: PLC0415 - optional backend

    reader = PdfReader(path)
    text = "\n".join((page.extract_text() or "") for page in reader.pages[:2])
    return text[:max_chars]


def docx_text(path: str, max_chars: int) -> str:
    import docx  # noqa: PLC0415 - optional backend

    collected: list[str] = []
    length = 0
    for paragraph in docx.Document(path).paragraphs:
        collected.append(paragraph.text)
        length += len(paragraph.text) + 1
        if length >= max_chars:  # enough to classify; stop walking the document
            break
    return "\n".join(collected)[:max_chars]


def xlsx_text(path: str, max_chars: int) -> str:
    import openpyxl  # noqa: PLC0415 - optional backend

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    cells: list[str] = []
    for row in sheet.iter_rows(max_row=20, values_only=True):
        cells += [str(c) for c in row if c is not None]
    return " ".join(cells)[:max_chars]


PARSERS = {"pdf": pdf_text, "docx": docx_text, "xlsx": xlsx_text}

#: The child's exit status when its parser raised.
EXIT_PARSE_FAILED = 3


def parse(kind: str, path: str, max_chars: int) -> str:
    """Text of ``path`` read by the ``kind`` parser; ``""`` for a kind with no parser.

    Raises:
        Exception: Whatever the parser raised on the file.
    """
    parser = PARSERS.get(kind)
    if parser is None:
        return ""  # the parent only asks for the kinds of extraction._LIBRARIES
    return parser(path, max_chars)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("usage: python -m cubby.adapters.parsers KIND PATH MAX_CHARS", file=sys.stderr)
        return 2
    kind, path, max_chars = argv
    try:
        text = parse(kind, path, int(max_chars))
    except Exception as exc:  # noqa: BLE001 - reported to the parent, which logs and counts it
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        return EXIT_PARSE_FAILED
    # Bytes, not text: the child's locale may not be UTF-8, and an accented
    # invoice must not turn into an encoding error.
    sys.stdout.buffer.write(text.encode("utf-8"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
